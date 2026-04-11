from __future__ import annotations

from collections import Counter
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal, InvalidOperation
from hashlib import sha256
from io import BytesIO
from typing import Any

from openpyxl import load_workbook


IMPORTANT_FIELDS = [
    "Id",
    "Name",
    "Temp Customer Status",
    "Sub. Status",
    "Source",
    "Sub-source",
    "Campaign Name",
    "Created Date",
    "Note",
    "Total Orders",
    "Total Revenue",
    "Route Name",
    "Payment Mode",
    "Due Since",
    "Last Payment Before Days",
    "Current Consumption",
    "Delivery Boy",
    "DND",
    "Last Delivery Date",
    "Area",
    "Mobile",
    "Hub",
    "Wallet Balance",
    "Effective Wallet Balance",
]


@dataclass
class ParsedWorkbook:
    file_name: str
    sheet_name: str
    fingerprint: str
    headers: list[str]
    row_count: int
    non_empty_counts: dict[str, int]
    status_samples: list[str]
    subscription_status_samples: list[str]
    rows: list[dict[str, Any]]


def _to_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _to_decimal(value: Any) -> float | None:
    if value in (None, "", "N/A"):
        return None
    try:
        return float(Decimal(str(value).replace(",", "").strip()))
    except (InvalidOperation, ValueError, TypeError):
        return None


def _to_int(value: Any) -> int | None:
    numeric = _to_decimal(value)
    if numeric is None:
        return None
    return int(numeric)


def _to_bool(value: Any) -> bool:
    normalized = _to_text(value).lower()
    return normalized in {"yes", "true", "1", "blocked"}


def _to_datetime(value: Any) -> datetime | None:
    if value in (None, "", "N/A"):
        return None
    if isinstance(value, datetime):
        return value
    raw = _to_text(value)
    if not raw:
        return None
    for fmt in (
        "%Y-%m-%d %H:%M:%S",
        "%d %b %Y %I:%M %p",
        "%d %b %Y",
        "%Y-%m-%d",
        "%d/%m/%Y",
    ):
        try:
            return datetime.strptime(raw, fmt)
        except ValueError:
            continue
    return None


def _row_hash(row: dict[str, Any]) -> str:
    parts = [
        _to_text(row.get("source_customer_id")),
        _to_text(row.get("mobile")),
        _to_text(row.get("name")),
        _to_text(row.get("temp_customer_status")),
        _to_text(row.get("subscription_status")),
        _to_text(row.get("wallet_balance_spend")),
        _to_text(row.get("effective_wallet_balance_current")),
    ]
    return sha256("|".join(parts).encode("utf-8")).hexdigest()


def parse_milkmaster_customer_workbook(file_name: str, content: bytes) -> ParsedWorkbook:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    sheet_name = workbook.sheetnames[0]
    worksheet = workbook[sheet_name]

    rows_iter = worksheet.iter_rows(values_only=True)
    headers = [_to_text(value) for value in next(rows_iter)]
    header_index = {header: idx for idx, header in enumerate(headers)}

    non_empty_counts: Counter[str] = Counter()
    status_counter: Counter[str] = Counter()
    subscription_counter: Counter[str] = Counter()
    normalized_rows: list[dict[str, Any]] = []

    def pick(row: tuple[Any, ...], column: str) -> Any:
        idx = header_index.get(column)
        if idx is None or idx >= len(row):
            return None
        return row[idx]

    for row in rows_iter:
        if row is None:
            continue

        normalized = {
            "source_customer_id": _to_text(pick(row, "Id")),
            "geo_location_html": _to_text(pick(row, "Geo location")),
            "alternate_mobile": _to_text(pick(row, "Alternate Mobile")),
            "mobile": _to_text(pick(row, "Mobile")),
            "name": _to_text(pick(row, "Name")),
            "temp_customer_status": _to_text(pick(row, "Temp Customer Status")),
            "subscription_status": _to_text(pick(row, "Sub. Status")),
            "follow_up_date": _to_datetime(pick(row, "Follow Up Date")),
            "source": _to_text(pick(row, "Source")),
            "sub_source": _to_text(pick(row, "Sub-source")),
            "campaign_name": _to_text(pick(row, "Campaign Name")),
            "created_date": _to_datetime(pick(row, "Created Date")),
            "note": _to_text(pick(row, "Note")),
            "total_orders": _to_int(pick(row, "Total Orders")),
            "total_revenue": _to_decimal(pick(row, "Total Revenue")),
            "route_name": _to_text(pick(row, "Route Name")),
            "payment_mode": _to_text(pick(row, "Payment Mode")),
            "first_delivery_date": _to_datetime(pick(row, "First Delivery Date")),
            "gst_number": _to_text(pick(row, "GST Number")),
            "due_since": _to_decimal(pick(row, "Due Since")),
            "last_payment_before_days": _to_int(pick(row, "Last Payment Before Days")),
            "current_consumption": _to_decimal(pick(row, "Current Consumption")),
            "delivery_preference": _to_text(pick(row, "Delivery Preferance")),
            "delivery_boy": _to_text(pick(row, "Delivery Boy")),
            "dnd": _to_bool(pick(row, "DND")),
            "created_by": _to_text(pick(row, "Created By")),
            "last_delivery_date": _to_datetime(pick(row, "Last Delivery Date")),
            "time_slot": _to_text(pick(row, "Time slot")),
            "area": _to_text(pick(row, "Area")),
            "email_id": _to_text(pick(row, "Email Id")),
            "is_blocked": _to_bool(pick(row, "Is Blocked")),
            "hub": _to_text(pick(row, "Hub")),
            "address": _to_text(pick(row, "Address")),
            "city": _to_text(pick(row, "City")),
            "crm_agent": _to_text(pick(row, "CRM Agent")),
            "sub_area": _to_text(pick(row, "Sub Area")),
            "wallet_balance_spend": _to_decimal(pick(row, "Wallet Balance")),
            "effective_wallet_balance_current": _to_decimal(pick(row, "Effective Wallet Balance")),
            "credit_limit": _to_decimal(pick(row, "Credit Limit")),
            "payment_type": _to_text(pick(row, "Payment Type")),
            "customer_type": _to_text(pick(row, "Customer Type")),
        }
        normalized["row_hash"] = _row_hash(normalized)
        normalized_rows.append(normalized)

        for field in IMPORTANT_FIELDS:
            value = pick(row, field)
            if value not in (None, "", " "):
                non_empty_counts[field] += 1

        if normalized["temp_customer_status"]:
            status_counter[normalized["temp_customer_status"]] += 1
        if normalized["subscription_status"]:
            subscription_counter[normalized["subscription_status"]] += 1

    fingerprint = sha256(content).hexdigest()
    return ParsedWorkbook(
        file_name=file_name,
        sheet_name=sheet_name,
        fingerprint=fingerprint,
        headers=headers,
        row_count=len(normalized_rows),
        non_empty_counts=dict(non_empty_counts),
        status_samples=[item for item, _count in status_counter.most_common(12)],
        subscription_status_samples=[item for item, _count in subscription_counter.most_common(12)],
        rows=normalized_rows,
    )
